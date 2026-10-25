
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE = 5 * 1024 * 1024


def extract_text_from_file(uploaded_file) -> str:
    """Extract text from a PDF, DOCX, or TXT upload."""

    if uploaded_file is None:
        raise ValueError("Please upload a file.")

    extension = Path(uploaded_file.name).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Upload a PDF, DOCX, or TXT file.")

    content = uploaded_file.getvalue()

    if not content:
        raise ValueError("The uploaded file is empty.")

    if len(content) > MAX_FILE_SIZE:
        raise ValueError("The file must be 5 MB or smaller.")

    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")

        elif extension == ".pdf":
            reader = PdfReader(BytesIO(content))

            if reader.is_encrypted:
                raise ValueError(
                    "Password-protected PDFs are not supported."
                )

            text = "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        else:
            document = Document(BytesIO(content))
            paragraphs = [
                paragraph.text for paragraph in document.paragraphs
            ]

            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

            text = "\n".join(paragraphs)

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            "Could not read the document. Check that it is valid "
            "and is not password-protected."
        ) from exc

    if not text.strip():
        raise ValueError(
            "No readable text found. Scanned PDFs may require OCR."
        )

    return text.strip()
